import streamlit as st
import google.generativeai as genai
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import io
import datetime

# ==========================================
# 1. SETUP & KONFIGURASI 
# ==========================================
st.set_page_config(page_title="LessonLab by Azaril", page_icon="🚀", layout="wide")

GOOGLE_API_KEY = st.secrets["GOOGLE_API_KEY"]
genai.configure(api_key=GOOGLE_API_KEY)
model = genai.GenerativeModel('gemini-2.5-flash')

# ==========================================
# 2. PANGKALAN DATA SILIBUS FIZIK KSSM (T4 & T5)
# ==========================================
silibus_fizik_kssm = {
    "Tingkatan 4": {
        "Bab 1: Pengukuran": ["1.1 Kuantiti Fizik", "1.2 Penyiasatan Saintifik"],
        "Bab 2: Daya dan Gerakan I": ["2.1 Gerakan Linear", "2.2 Graf Gerakan Linear", "2.3 Gerakan Jatuh Bebas", "2.4 Inersia", "2.5 Momentum", "2.6 Daya", "2.7 Impuls dan Daya Impuls", "2.8 Berat"],
        "Bab 3: Kegravitian": ["3.1 Hukum Kegravitian Semesta Newton", "3.2 Hukum Kepler", "3.3 Satelit Buatan Manusia"],
        "Bab 4: Haba": ["4.1 Keseimbangan Terma", "4.2 Muatan Haba Tentu", "4.3 Haba Pendam Tentu", "4.4 Hukum Gas"],
        "Bab 5: Gelombang": ["5.1 Asas Gelombang", "5.2 Pelembapan dan Resonans", "5.3 Pantulan Gelombang", "5.4 Pembiasan Gelombang", "5.5 Pembelauan Gelombang", "5.6 Interferens Gelombang", "5.7 Gelombang Elektromagnet"],
        "Bab 6: Cahaya dan Optik": ["6.1 Pembiasan Cahaya", "6.2 Pantulan Dalam Penuh", "6.3 Pembentukan Imej oleh Kanta", "6.4 Formula Kanta Nipis", "6.5 Peralatan Optik", "6.6 Pembentukan Imej oleh Cermin Sfera"]
    },
    "Tingkatan 5": {
        "Bab 1: Daya dan Gerakan II": ["1.1 Daya Paduan", "1.2 Leraian Daya", "1.3 Keseimbangan Daya", "1.4 Kekenyalan"],
        "Bab 2: Tekanan": ["2.1 Tekanan Cecair", "2.2 Tekanan Atmosfera", "2.3 Tekanan Gas", "2.4 Prinsip Pascal", "2.5 Prinsip Archimedes", "2.6 Prinsip Bernoulli"],
        "Bab 3: Keelektrikan": ["3.1 Arus dan Beza Keupayaan", "3.2 Rintangan", "3.3 Daya Gerak Elektrik (d.g.e) dan Rintangan Dalam", "3.4 Tenaga dan Kuasa Elektrik"],
        "Bab 4: Keelektromagnetan": ["4.1 Daya ke atas Konduktor Pembawa Arus dalam Medan Magnet", "4.2 Aruhan Elektromagnet", "4.3 Transformer"],
        "Bab 5: Elektronik": ["5.1 Elektron", "5.2 Diod Semikonduktor", "5.3 Transistor"],
        "Bab 6: Fizik Nuklear": ["6.1 Reputan Radioaktif", "6.2 Tenaga Nuklear"],
        "Bab 7: Fizik Kuantum": ["7.1 Teori Kuantum Cahaya", "7.2 Kesan Fotoelektrik", "7.3 Teori Fotoelektrik Einstein"]
    }
}

# ==========================================
# 3. SIDEBAR (PILIHAN MOD APLIKASI)
# ==========================================
st.sidebar.image("https://cdn-icons-png.flaticon.com/512/3163/3163288.png", width=80)
st.sidebar.title("Menu Utama")
pilihan_mod = st.sidebar.radio("Pilih Jenis RPH:", ["📝 Ulangkaji SPM", "🚀 PdP 5E"])
st.sidebar.divider()
st.sidebar.info("LessonLab by Azaril. Lesson design system")


# =========================================================================================
# MOD 1: RPH ULANGKAJI SPM
# =========================================================================================
if pilihan_mod == "📝 Ulangkaji SPM":
    
    st.title("📝 LessonLab by Azaril Ulangkaji SPM")
    st.markdown("Isi maklumat di bawah. Anda kini boleh memilih **Tingkatan 4 & 5 sekaligus** untuk kelas ulangkaji bersepadu.")

    kolom_kiri, kolom_kanan = st.columns(2)

    with kolom_kiri:
        st.subheader("📅 Maklumat Kelas")
        tarikh_pilihan = st.date_input("Tarikh Kelas", key="tarikh_uk")
        hari_dict = {0: "Isnin", 1: "Selasa", 2: "Rabu", 3: "Khamis", 4: "Jumaat", 5: "Sabtu", 6: "Ahad"}
        hari_auto = hari_dict[tarikh_pilihan.weekday()]
        
        col_m1, col_m2 = st.columns(2)
        with col_m1: masa_mula = st.time_input("Masa Mula", value=datetime.time(7, 40), step=600, key="mula_uk")
        with col_m2: masa_tamat = st.time_input("Masa Tamat", value=datetime.time(9, 00), step=600, key="tamat_uk")

        t1 = datetime.datetime.combine(datetime.date.today(), masa_mula)
        t2 = datetime.datetime.combine(datetime.date.today(), masa_tamat)
        tempoh_minit = int((t2 - t1).total_seconds() / 60)
        if tempoh_minit < 0: tempoh_minit += 1440 
        
        col_k1, col_k2 = st.columns(2)
        with col_k1: nama_kelas = st.text_input("Nama Kelas", "5 Sains 1", key="kelas_uk")
        with col_k2: bil_murid = st.number_input("Bilangan Murid", min_value=1, max_value=50, value=30, key="murid_uk")
        
        st.info(f"📌 **Rumusan:** {hari_auto} | Tempoh: {tempoh_minit} Minit")

    with kolom_kanan:
        st.subheader("🎯 Fokus Ulangkaji & Topik")
        format_kertas = st.selectbox("Fokus Kertas SPM", [
            "Kertas 1 (Objektif)", 
            "Kertas 2 - Bahagian A (Struktur)", 
            "Kertas 2 - Bahagian B & C (Esei)", 
            "Kertas 3 (Ujian Amali Bersepadu)"
        ])
        bahan_bbm = st.text_input("Bahan Ulangkaji / Modul", value="Modul Latih Tubi SPM")

        tingkatan_pilihan = st.multiselect(
            "Pilih Tingkatan (Boleh pilih kedua-duanya)", 
            ["Tingkatan 4", "Tingkatan 5"], 
            default=["Tingkatan 5"]
        )
        
        senarai_bab_gabungan = []
        for t in tingkatan_pilihan:
            senarai_bab_gabungan.extend(list(silibus_fizik_kssm[t].keys()))
        
        pilih_semua_bab = st.checkbox("☑️ Pilih Semua Bab")
        if pilih_semua_bab:
            bab_dipilih = st.multiselect("Pilih Bab", options=senarai_bab_gabungan, default=senarai_bab_gabungan)
        else:
            bab_dipilih = st.multiselect("Pilih Bab", options=senarai_bab_gabungan)
        
        senarai_sub_gabungan = []
        for t in tingkatan_pilihan:
            for b in bab_dipilih:
                if b in silibus_fizik_kssm[t]:
                    senarai_sub_gabungan.extend(silibus_fizik_kssm[t][b])
                    
        pilih_semua_sub = st.checkbox("☑️ Pilih Semua Sub-topik (Pilihan)")
        if pilih_semua_sub:
            sub_bab_dipilih = st.multiselect("Pilih Sub-topik Fokus (Optional)", options=senarai_sub_gabungan, default=senarai_sub_gabungan)
        else:
            sub_bab_dipilih = st.multiselect("Pilih Sub-topik Fokus (Optional)", options=senarai_sub_gabungan)
        
        txt_tingkatan = " & ".join(tingkatan_pilihan)
        
        if pilih_semua_bab:
            topik_penuh = "Semua Bab"
        else:
            topik_penuh = ", ".join(bab_dipilih) if not sub_bab_dipilih else f"{', '.join(bab_dipilih)} ({', '.join(sub_bab_dipilih)})"

    st.divider()

    st.subheader("⏳ Agihan Masa (Jumlah Mesti = Masa Kelas)")
    col1, col2, col3, col4, col5 = st.columns(5)
    with col1: m_recall = st.number_input("Recall", value=10)
    with col2: m_teknik = st.number_input("Teknik Menjawab", value=20)
    with col3: m_latih = st.number_input("Latih Tubi", value=30)
    with col4: m_bincang = st.number_input("Perbincangan", value=20)
    with col5: m_tutup = st.number_input("Penutup", value=10)

    jumlah_masa_fasa = m_recall + m_teknik + m_latih + m_bincang + m_tutup
    st.caption(f"**Jumlah Masa:** {jumlah_masa_fasa} Minit / {tempoh_minit} Minit")

    def jana_rph_ai_ulangkaji(topik, ting, bil_murid, format_spm, bbm):
        prompt = f"""
        Bertindak sebagai Guru Cemerlang Fizik SPM.
        Bina RPH Ulangkaji untuk topik: {topik} ({ting}).
        Fokus: Latih Tubi {format_spm}.
        Bahan: {bbm}.
        Bilangan murid: {bil_murid} orang.

        AMARAN PENTING: DILARANG MENULIS SEBARANG NOMBOR MASA ATAU MINIT dalam teks penerangan anda.
        Terus nyatakan langkah-langkah aktiviti sahaja.

        Format WAJIB:
        [OBJEKTIF] (Tulis 2 objektif berfokus teknik menjawab SPM)
        ---
        [KRITERIA KEJAYAAN] (Tulis 2-3 kriteria kejayaan spesifik)
        ---
        [ABM] (Senarai bahan: {bbm})
        ---
        [EMK] (Elemen merentas kurikulum sesuai)
        ---
        [KBAT] (Kata tugas soalan peperiksaan)
        ---
        [RECALL] (Aktiviti mengingat semula formula/konsep Fizik berkaitan. Tanpa minit)
        ---
        [TEKNIK] (Penerangan guru tentang cara dapat markah penuh {format_spm}. Tanpa minit)
        ---
        [LATIH_TUBI] (Aktiviti murid menjawab soalan secara kendiri. Tanpa minit)
        ---
        [PERBINCANGAN] (Aktiviti menyemak jawapan dan pendedahan skema markah. Tanpa minit)
        ---
        [PENUTUP] (Rumusan prestasi kelas. Tanpa minit)
        ---
        [REFLEKSI] (Refleksi guru realistik selepas kelas selesai).
        """
        response = model.generate_content(prompt)
        return response.text

    def bina_excel_ulangkaji(data_ai, data_input):
        teks_bersih = data_ai.replace("**", "")
        def ekstrak(tag, teks):
            if tag in teks:
                mula = teks.split(tag)[1]
                if "---" in mula:
                    return mula.split("---")[0].strip()
                else:
                    return mula.split("[")[0].strip()
            return ""

        excel_data = {
            'objektif': ekstrak("[OBJEKTIF]", teks_bersih), 'kk': ekstrak("[KRITERIA KEJAYAAN]", teks_bersih),
            'abm': ekstrak("[ABM]", teks_bersih), 'emk': ekstrak("[EMK]", teks_bersih), 
            'kbat': ekstrak("[KBAT]", teks_bersih), 'recall': ekstrak("[RECALL]", teks_bersih), 
            'teknik': ekstrak("[TEKNIK]", teks_bersih), 'latih': ekstrak("[LATIH_TUBI]", teks_bersih),
            'bincang': ekstrak("[PERBINCANGAN]", teks_bersih), 'tutup': ekstrak("[PENUTUP]", teks_bersih),
            'refleksi': ekstrak("[REFLEKSI]", teks_bersih)
        }

        wb = openpyxl.Workbook()
        ws = wb.active; ws.title = "RPH_Ulangkaji"

        navy_dark, blue_accent = "1B365D", "E6F0FA"
        fill_title = PatternFill(start_color=navy_dark, end_color=navy_dark, fill_type="solid")
        fill_section = PatternFill(start_color="2C5282", end_color="2C5282", fill_type="solid")
        fill_light = PatternFill(start_color=blue_accent, end_color=blue_accent, fill_type="solid")
        f_title = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
        f_sec = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        f_b = Font(name="Segoe UI", size=11, bold=True); f_norm = Font(name="Segoe UI", size=11)
        border = Border(left=Side(style='thin', color='A0A0A0'), right=Side(style='thin', color='A0A0A0'), top=Side(style='thin', color='A0A0A0'), bottom=Side(style='thin', color='A0A0A0'))
        al_c = Alignment(horizontal="center", vertical="center", wrap_text=True); al_l = Alignment(horizontal="left", vertical="center", wrap_text=True); al_lt = Alignment(horizontal="left", vertical="top", wrap_text=True)

        def gayakan(r, c, font, align, fill=None):
            cell = ws.cell(row=r, column=c)
            cell.font, cell.alignment, cell.border = font, align, border
            if fill: cell.fill = fill

        ws.merge_cells("A1:D1"); ws["A1"] = "RANCANGAN PENGAJARAN HARIAN (ULANGKAJI)"
        gayakan(1, 1, f_title, al_c, fill_title); ws.row_dimensions[1].height = 25
        ws.merge_cells("A2:D2"); ws["A2"] = f"FIZIK KSSM | {data_input['tingkatan'].upper()}"
        gayakan(2, 1, Font(name="Segoe UI", size=11, bold=True, color="FFFFFF"), al_c, fill_title)
        
        ws.merge_cells("A4:D4"); ws["A4"] = " MAKLUMAT KELAS"
        gayakan(4, 1, f_sec, al_l, fill_section)
        ws["A5"], ws["B5"] = "Hari & Tarikh:", f"{data_input['hari']}, {data_input['tarikh']}"
        ws["C5"], ws["D5"] = "Masa & Tempoh:", f"{data_input['masa']} ({data_input['tempoh']} min)"
        ws["A6"], ws["B6"] = "Kelas:", data_input['kelas']
        ws["C6"], ws["D6"] = "Kehadiran:", f"_____ / {data_input['bil_murid']} orang"
        for r in [5,6]:
            gayakan(r, 1, f_b, al_l, fill_light); gayakan(r, 2, f_norm, al_l)
            gayakan(r, 3, f_b, al_l, fill_light); gayakan(r, 4, f_norm, al_l)

        ws.merge_cells("A8:D8"); ws["A8"] = " MAKLUMAT FOKUS ULANGKAJI"
        gayakan(8, 1, f_sec, al_l, fill_section)
        
        pdp_items = [("Topik/Bab", data_input['topik']), ("Fokus Kertas", data_input['format_spm']),
                     ("Objektif", excel_data['objektif']), ("Kriteria", excel_data['kk'])]
        r = 9
        for lbl, val in pdp_items:
            ws[f"A{r}"] = lbl; ws.merge_cells(f"B{r}:D{r}"); ws[f"B{r}"] = val
            gayakan(r, 1, f_b, al_lt, fill_light); 
            for c in [2,3,4]: gayakan(r, c, f_norm, al_lt)
            ws.row_dimensions[r].height = max(40, (len(str(val)) // 80 + 1) * 18)
            r+=1

        r+=1; ws.merge_cells(f"A{r}:D{r}"); ws[f"A{r}"] = " STRATEGI ULANGKAJI"
        gayakan(r, 1, f_sec, al_l, fill_section); r+=1
        
        # Penutup diletakkan bersama refleksi
        aktiviti = [
            (f"RECALL\n{data_input['m_recall']} Min", excel_data['recall']),
            (f"TEKNIK\n{data_input['m_teknik']} Min", excel_data['teknik']),
            (f"LATIH TUBI\n{data_input['m_latih']} Min", excel_data['latih']),
            (f"BINCANG\n{data_input['m_bincang']} Min", excel_data['bincang'])
        ]
        for lbl, val in aktiviti:
            ws[f"A{r}"] = lbl; ws.merge_cells(f"B{r}:D{r}"); ws[f"B{r}"] = val
            gayakan(r, 1, f_b, Alignment(horizontal="center", vertical="top", wrap_text=True), fill_light)
            for c in [2,3,4]: gayakan(r, c, f_norm, al_lt)
            ws.row_dimensions[r].height = max(60, (len(str(val)) // 80 + 2) * 18)
            r+=1

        r+=1; ws.merge_cells(f"A{r}:D{r}"); ws[f"A{r}"] = " PENUTUP & REFLEKSI"
        gayakan(r, 1, f_sec, al_l, fill_section); r+=1
        
        ws[f"A{r}"] = f"Rumusan\n({data_input['m_tutup']} Min)"; ws.merge_cells(f"B{r}:D{r}"); ws[f"B{r}"] = excel_data['tutup']
        gayakan(r, 1, f_b, al_lt, fill_light)
        for c in [2,3,4]: gayakan(r, c, f_norm, al_lt)
        ws.row_dimensions[r].height = max(45, (len(str(excel_data['tutup'])) // 80 + 2) * 18); r+=1
        
        ws[f"A{r}"] = "Refleksi Guru"; ws.merge_cells(f"B{r}:D{r}"); ws[f"B{r}"] = excel_data['refleksi']
        gayakan(r, 1, f_b, al_lt, fill_light)
        for c in [2,3,4]: gayakan(r, c, f_norm, al_lt)
        ws.row_dimensions[r].height = max(60, (len(str(excel_data['refleksi'])) // 80 + 2) * 18)

        ws.column_dimensions['A'].width = 18; ws.column_dimensions['B'].width = 25
        ws.column_dimensions['C'].width = 25; ws.column_dimensions['D'].width = 25
        
        output = io.BytesIO(); wb.save(output)
        return output.getvalue()

    if st.button("🚀 Jana RPH Ulangkaji SPM Sekarang", use_container_width=True, type="primary"):
        if not bab_dipilih:
            st.warning("Sila pilih sekurang-kurangnya satu Bab.")
        elif jumlah_masa_fasa != tempoh_minit:
            st.error(f"Jumlah masa fasa ({jumlah_masa_fasa} minit) tidak sama dengan tempoh kelas ({tempoh_minit} minit).")
        else:
            with st.spinner("🧠 LessonLab sedang menyusun strategi ulangkaji..."):
                hasil = jana_rph_ai_ulangkaji(topik_penuh, txt_tingkatan, bil_murid, format_kertas, bahan_bbm)
                data_input = {
                    "tingkatan": txt_tingkatan, "topik": topik_penuh, "bil_murid": bil_murid,
                    "hari": hari_auto, "tarikh": tarikh_pilihan.strftime('%d/%m/%Y'),
                    "masa": f"{masa_mula.strftime('%I:%M %p')} - {masa_tamat.strftime('%I:%M %p')}",
                    "tempoh": tempoh_minit, "kelas": nama_kelas, "format_spm": format_kertas, "bbm": bahan_bbm,
                    "m_recall": m_recall, "m_teknik": m_teknik, "m_latih": m_latih,
                    "m_bincang": m_bincang, "m_tutup": m_tutup
                }
                file_excel = bina_excel_ulangkaji(hasil, data_input)
                st.success("🎉 RPH Berhasil Dijana!")
                st.download_button(
                    label="📥 Muat Turun RPH Ulangkaji (Excel)",
                    data=file_excel, file_name=f"RPH_Ulangkaji_Fizik_{nama_kelas}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )

# =========================================================================================
# MOD 2: RPH FIZIK 5E KSSM (PDP BIASA)
# =========================================================================================
elif pilihan_mod == "🚀 PdP 5E":
    
    st.title("🚀 LessonLab by Azaril PdP 5E")
    st.markdown("Isi maklumat automatik di bawah dan biarkan LessonLab hasilkan RPH anda.")

    kolom_kiri, kolom_kanan = st.columns(2)

    with kolom_kiri:
        st.subheader("📅 Maklumat Kelas")
        tarikh_pilihan = st.date_input("Tarikh Kelas", key="tarikh_5e")
        hari_dict = {0: "Isnin", 1: "Selasa", 2: "Rabu", 3: "Khamis", 4: "Jumaat", 5: "Sabtu", 6: "Ahad"}
        hari_auto = hari_dict[tarikh_pilihan.weekday()]
        
        col_m1, col_m2 = st.columns(2)
        with col_m1: masa_mula = st.time_input("Masa Mula", value=datetime.time(7, 40), step=600, key="mula_5e")
        with col_m2: masa_tamat = st.time_input("Masa Tamat", value=datetime.time(9, 40), step=600, key="tamat_5e")

        t1 = datetime.datetime.combine(datetime.date.today(), masa_mula)
        t2 = datetime.datetime.combine(datetime.date.today(), masa_tamat)
        tempoh_minit = int((t2 - t1).total_seconds() / 60)
        if tempoh_minit < 0: tempoh_minit += 1440 
        
        col_k1, col_k2 = st.columns(2)
        with col_k1: nama_kelas = st.text_input("Nama Kelas", "5 Sains 1", key="kelas_5e")
        with col_k2: bil_murid = st.number_input("Bilangan Murid", min_value=1, max_value=50, value=30, key="murid_5e")
        
        st.info(f"📌 **Rumusan:** Hari {hari_auto}, {tarikh_pilihan.strftime('%d/%m/%Y')} | Tempoh: {tempoh_minit} Minit")

    with kolom_kanan:
        st.subheader("📚 Maklumat Topik (Silibus Auto)")
        tingkatan = st.selectbox("Pilih Tingkatan", ["Tingkatan 4", "Tingkatan 5"], key="ting_5e")
        
        senarai_bab = list(silibus_fizik_kssm[tingkatan].keys())
        bab = st.selectbox("Pilih Bab Utama", senarai_bab, key="bab_5e")
        
        senarai_sub = silibus_fizik_kssm[tingkatan][bab]
        sub_bab = st.multiselect("Pilih Sub-Topik (Boleh pilih lebih dari 1)", senarai_sub, key="sub_5e")
        
        if sub_bab:
            topik_penuh = f"{bab} ({', '.join(sub_bab)})"
        else:
            topik_penuh = bab

    st.divider()

    st.subheader("⏳ Agihan Masa 5E (Pastikan Jumlah = Masa Kelas)")
    col1, col2, col3, col4, col5 = st.columns(5)
    with col1: m_engage = st.number_input("Engage", value=5, min_value=1)
    with col2: m_explore = st.number_input("Explore", value=min(20, tempoh_minit//4), min_value=1)
    with col3: m_explain = st.number_input("Explain", value=min(20, tempoh_minit//4), min_value=1)
    with col4: m_elaborate = st.number_input("Elaborate", value=min(10, tempoh_minit//5), min_value=1)
    with col5: m_evaluate = st.number_input("Evaluate", value=5, min_value=1)

    jumlah_masa_5e = m_engage + m_explore + m_explain + m_elaborate + m_evaluate
    st.caption(f"**Jumlah Masa Disusun:** {jumlah_masa_5e} Minit daripada {tempoh_minit} Minit")

    def jana_rph_ai_5e(topik, ting, bil_murid):
        prompt = f"""
        Bertindak sebagai Guru Cemerlang Fizik KSSM.
        Bina Rancangan Pengajaran Harian (RPH) untuk topik: {topik} ({ting}).
        Bilangan murid: {bil_murid} orang.

        AMARAN PENTING: DILARANG MENULIS SEBARANG NOMBOR MASA ATAU MINIT dalam teks penerangan anda.
        Terus nyatakan langkah-langkah aktiviti sahaja.

        Format WAJIB:
        [OBJEKTIF] (Tulis 2 objektif pembelajaran)
        ---
        [KRITERIA KEJAYAAN] (Tulis 2 atau 3 kriteria kejayaan spesifik)
        ---
        [ABM] (Senarai bahan bantu mengajar)
        ---
        [EMK] (Elemen merentas kurikulum yang sesuai)
        ---
        [KBAT] (Kemahiran berfikir aras tinggi)
        ---
        [ENGAGE] (Aktiviti set induksi menarik. Tanpa minit)
        ---
        [EXPLORE] (Aktiviti penerokaan berpusatkan murid. Tanpa minit)
        ---
        [EXPLAIN] (Penerangan konsep oleh guru. Tanpa minit)
        ---
        [ELABORATE] (Aktiviti pengembangan/aplikasi. Tanpa minit)
        ---
        [EVALUATE] (Kaedah penilaian/pentaksiran formatif. Tanpa minit)
        ---
        [RUMUSAN] (Rumusan penutup guru. Tanpa minit)
        ---
        [REFLEKSI] (Reka satu perenggan refleksi guru yang realistik).
        """
        response = model.generate_content(prompt)
        return response.text

    def bina_excel_5e(data_ai, data_input):
        teks_bersih = data_ai.replace("**", "")
        def ekstrak(tag, teks):
            if tag in teks:
                mula = teks.split(tag)[1]
                if "---" in mula:
                    return mula.split("---")[0].strip()
                else:
                    return mula.split("[")[0].strip() 
            return ""

        excel_data = {
            'objektif': ekstrak("[OBJEKTIF]", teks_bersih), 'kk': ekstrak("[KRITERIA KEJAYAAN]", teks_bersih),
            'abm': ekstrak("[ABM]", teks_bersih), 'emk': ekstrak("[EMK]", teks_bersih), 
            'kbat': ekstrak("[KBAT]", teks_bersih), 'engage': ekstrak("[ENGAGE]", teks_bersih), 
            'explore': ekstrak("[EXPLORE]", teks_bersih), 'explain': ekstrak("[EXPLAIN]", teks_bersih),
            'elaborate': ekstrak("[ELABORATE]", teks_bersih), 'evaluate': ekstrak("[EVALUATE]", teks_bersih),
            'rumusan': ekstrak("[RUMUSAN]", teks_bersih), 'refleksi': ekstrak("[REFLEKSI]", teks_bersih)
        }

        wb = openpyxl.Workbook()
        ws = wb.active; ws.title = "RPH_5E"

        navy_dark, blue_accent = "1B365D", "E6F0FA"
        fill_title = PatternFill(start_color=navy_dark, end_color=navy_dark, fill_type="solid")
        fill_section = PatternFill(start_color="2C5282", end_color="2C5282", fill_type="solid")
        fill_light = PatternFill(start_color=blue_accent, end_color=blue_accent, fill_type="solid")
        f_title = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
        f_sec = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        f_b = Font(name="Segoe UI", size=11, bold=True); f_norm = Font(name="Segoe UI", size=11)
        border = Border(left=Side(style='thin', color='A0A0A0'), right=Side(style='thin', color='A0A0A0'), top=Side(style='thin', color='A0A0A0'), bottom=Side(style='thin', color='A0A0A0'))
        al_c = Alignment(horizontal="center", vertical="center", wrap_text=True); al_l = Alignment(horizontal="left", vertical="center", wrap_text=True); al_lt = Alignment(horizontal="left", vertical="top", wrap_text=True)

        def gayakan(r, c, font, align, bord=border, fill=None):
            cell = ws.cell(row=r, column=c)
            cell.font, cell.alignment = font, align
            if bord: cell.border = bord
            if fill: cell.fill = fill

        ws.merge_cells("A1:D1"); ws["A1"] = "RANCANGAN PENGAJARAN HARIAN (RPH 5E)"
        gayakan(1, 1, f_title, al_c, None, fill_title); ws.row_dimensions[1].height = 25
        ws.merge_cells("A2:D2"); ws["A2"] = f"FIZIK | KSSM | {data_input['tingkatan'].upper()}"
        gayakan(2, 1, Font(name="Segoe UI", size=11, bold=True, color="FFFFFF"), al_c, None, fill_title)
        
        ws.merge_cells("A4:D4"); ws["A4"] = " MAKLUMAT KELAS"
        gayakan(4, 1, f_sec, al_l, border, fill_section)
        ws["A5"], ws["B5"] = "Hari & Tarikh:", f"{data_input['hari']}, {data_input['tarikh']}"
        ws["C5"], ws["D5"] = "Masa & Tempoh:", f"{data_input['masa']} ({data_input['tempoh']} minit)"
        ws["A6"], ws["B6"] = "Kelas:", data_input['kelas']
        ws["C6"], ws["D6"] = "Kehadiran:", f"_____ / {data_input['bil_murid']} orang"
        for r in [5,6]:
            gayakan(r, 1, f_b, al_l, border, fill_light); gayakan(r, 2, f_norm, al_l, border)
            gayakan(r, 3, f_b, al_l, border, fill_light); gayakan(r, 4, f_norm, al_l, border)

        ws.merge_cells("A8:D8"); ws["A8"] = " MAKLUMAT PENGAJARAN & PEMBELAJARAN"
        gayakan(8, 1, f_sec, al_l, border, fill_section)
        
        pdp_items = [("Topik", data_input['topik']), ("Objektif", excel_data['objektif']), 
                     ("Kriteria Kejayaan", excel_data['kk']), ("BBM", excel_data['abm']),
                     ("EMK", excel_data['emk']), ("KBAT", excel_data['kbat'])]
        r = 9
        for lbl, val in pdp_items:
            ws[f"A{r}"] = lbl; ws.merge_cells(f"B{r}:D{r}"); ws[f"B{r}"] = val
            gayakan(r, 1, f_b, al_lt, border, fill_light)
            for c in [2,3,4]: gayakan(r, c, f_norm, al_lt, border)
            ws.row_dimensions[r].height = max(25, (len(str(val)) // 80 + 1) * 18)
            r+=1

        r+=1; ws.merge_cells(f"A{r}:D{r}"); ws[f"A{r}"] = " AKTIVITI PENGAJARAN (MODEL 5E)"
        gayakan(r, 1, f_sec, al_l, border, fill_section); r+=1
        
        aktiviti = [
            (f"ENGAGE\n{data_input['m_engage']} Min", excel_data['engage']),
            (f"EXPLORE\n{data_input['m_explore']} Min", excel_data['explore']),
            (f"EXPLAIN\n{data_input['m_explain']} Min", excel_data['explain']),
            (f"ELABORATE\n{data_input['m_elaborate']} Min", excel_data['elaborate']),
            (f"EVALUATE\n{data_input['m_evaluate']} Min", excel_data['evaluate'])
        ]
        for lbl, val in aktiviti:
            ws[f"A{r}"] = lbl; ws.merge_cells(f"B{r}:D{r}"); ws[f"B{r}"] = val
            gayakan(r, 1, f_b, Alignment(horizontal="center", vertical="top", wrap_text=True), border, fill_light)
            for c in [2,3,4]: gayakan(r, c, f_norm, al_lt, border)
            ws.row_dimensions[r].height = max(45, (len(str(val)) // 80 + 2) * 18)
            r+=1

        r+=1; ws.merge_cells(f"A{r}:D{r}"); ws[f"A{r}"] = " PENUTUP & REFLEKSI"
        gayakan(r, 1, f_sec, al_l, border, fill_section); r+=1
        
        ws[f"A{r}"] = "Rumusan"; ws.merge_cells(f"B{r}:D{r}"); ws[f"B{r}"] = excel_data['rumusan']
        gayakan(r, 1, f_b, al_lt, border, fill_light)
        for c in [2,3,4]: gayakan(r, c, f_norm, al_lt, border)
        ws.row_dimensions[r].height = max(45, (len(str(excel_data['rumusan'])) // 80 + 2) * 18); r+=1
        
        ws[f"A{r}"] = "Refleksi Guru"; ws.merge_cells(f"B{r}:D{r}"); ws[f"B{r}"] = excel_data['refleksi']
        gayakan(r, 1, f_b, al_lt, border, fill_light)
        for c in [2,3,4]: gayakan(r, c, f_norm, al_lt, border)
        ws.row_dimensions[r].height = max(60, (len(str(excel_data['refleksi'])) // 80 + 2) * 18)
        
        ws.column_dimensions['A'].width = 20; ws.column_dimensions['B'].width = 25
        ws.column_dimensions['C'].width = 25; ws.column_dimensions['D'].width = 30
        
        output = io.BytesIO(); wb.save(output)
        return output.getvalue()

    if st.button("🚀 Jana RPH 5E Sekarang", use_container_width=True, type="primary"):
        if not sub_bab:
            st.warning("Sila pilih sekurang-kurangnya satu Sub-Topik.")
        elif jumlah_masa_5e != tempoh_minit:
            st.error(f"Jumlah masa fasa 5E ({jumlah_masa_5e} minit) tidak sama dengan tempoh kelas ({tempoh_minit} minit). Sila betulkan.")
        else:
            with st.spinner("🧠 LessonLab sedang menyusun RPH..."):
                hasil = jana_rph_ai_5e(topik_penuh, tingkatan, bil_murid)
                data_input = {
                    "tingkatan": tingkatan, "topik": topik_penuh, "bil_murid": bil_murid,
                    "hari": hari_auto, "tarikh": tarikh_pilihan.strftime('%d/%m/%Y'),
                    "masa": f"{masa_mula.strftime('%I:%M %p')} - {masa_tamat.strftime('%I:%M %p')}",
                    "tempoh": tempoh_minit, "kelas": nama_kelas,
                    "m_engage": m_engage, "m_explore": m_explore, "m_explain": m_explain,
                    "m_elaborate": m_elaborate, "m_evaluate": m_evaluate
                }
                file_excel = bina_excel_5e(hasil, data_input)
                st.success("🎉 RPH Berjaya Dijana!")
                st.download_button(
                    label="📥 Muat Turun RPH 5E (Excel)",
                    data=file_excel, file_name=f"RPH_5E_{nama_kelas}_{tarikh_pilihan}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary"
                )
